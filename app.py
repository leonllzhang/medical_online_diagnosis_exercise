#!/usr/bin/env python3
"""互联网诊疗医务人员准入习题考试系统"""
import json
import os
from datetime import datetime, timedelta, timezone

import random as _random

from flask import (
    Flask,
    abort,
    jsonify,
    render_template,
    request,
    send_file,
)
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from config import Config
from models import db, Examinee, ExamPaper, ExamSession, Role, RoleConfig

app = Flask(__name__)
app.config.from_object(Config)
db.init_app(app)


# ---------------------------------------------------------------------------
# 辅助函数
# ---------------------------------------------------------------------------

def load_questions_from_db_or_file():
    """加载题库：优先从数据库试卷表加载，否则从静态文件加载"""
    paper = ExamPaper.query.order_by(ExamPaper.id.desc()).first()
    if paper:
        return json.loads(paper.questions_json)
    # 回退到静态文件
    q_path = os.path.join(Config.BASE_DIR, "static", "js", "questions.json")
    if os.path.exists(q_path):
        with open(q_path, "r", encoding="utf-8") as f:
            return json.load(f)
    return []


# ---------------------------------------------------------------------------
# 页面路由
# ---------------------------------------------------------------------------

@app.route("/")
def exam_page():
    """考试页面 - 扫码后进入"""
    return render_template("exam.html")


@app.route("/admin")
def admin_page():
    """管理后台页面"""
    return render_template("admin.html")


# ---------------------------------------------------------------------------
# API - 题库
# ---------------------------------------------------------------------------

@app.route("/api/questions", methods=["GET"])
def get_questions():
    """获取所有题目（不返回答案，用于前端展示）"""
    questions = load_questions_from_db_or_file()
    safe = []
    for q in questions:
        safe.append(
            {
                "id": q["id"],
                "chapterId": q.get("chapterId", 0),
                "chapterTitle": q.get("chapterTitle", ""),
                "type": q["type"],
                "stem": q["stem"],
                "options": q["options"],
            }
        )
    return jsonify({"code": 0, "data": safe, "total": len(safe)})


# ---------------------------------------------------------------------------
# API - 提交答题结果
# ---------------------------------------------------------------------------

@app.route("/api/submit", methods=["POST"])
def submit_exam():
    """提交答题结果"""
    data = request.get_json(force=True)
    name = (data.get("name") or "").strip()
    department = (data.get("department") or "").strip()
    answers = data.get("answers")  # [{questionId, selected}, ...]

    if not name or not department:
        return jsonify({"code": 1, "msg": "请填写姓名和科室"}), 400
    if not answers or not isinstance(answers, list):
        return jsonify({"code": 1, "msg": "答题数据无效"}), 400

    questions = load_questions_from_db_or_file()
    q_map = {q["id"]: q for q in questions}

    # 判题
    correct_count = 0
    detail = []
    for ans in answers:
        qid = ans.get("questionId")
        selected = ans.get("selected", [])
        q = q_map.get(qid)
        if not q:
            continue
        correct_set = set(q["answer"])
        selected_set = set(selected)
        is_correct = correct_set == selected_set
        if is_correct:
            correct_count += 1
        detail.append(
            {
                "questionId": qid,
                "selected": selected,
                "correct": is_correct,
                "correctAnswer": q["answer"],
            }
        )

    total = len([a for a in answers if a.get("questionId") in q_map])
    percentage = round(correct_count / total * 100, 1) if total > 0 else 0
    status = "pass" if percentage >= app.config["PASS_THRESHOLD"] else "fail"

    record = Examinee(
        name=name,
        department=department,
        score=correct_count,
        total=total,
        percentage=percentage,
        status=status,
        answers_json=json.dumps(detail, ensure_ascii=False),
        ip_address=request.remote_addr or "",
    )
    db.session.add(record)
    db.session.commit()

    return jsonify(
        {
            "code": 0,
            "data": {
                "id": record.id,
                "score": correct_count,
                "total": total,
                "percentage": percentage,
                "status": status,
            },
        }
    )


# ---------------------------------------------------------------------------
# API - 考核记录查询
# ---------------------------------------------------------------------------

@app.route("/api/records", methods=["GET"])
def get_records():
    """获取考核记录列表（分页 + 筛选）"""
    page = request.args.get("page", 1, type=int)
    per_page = request.args.get("per_page", 20, type=int)
    per_page = min(per_page, 200)
    name = request.args.get("name", "").strip()
    department = request.args.get("department", "").strip()
    status = request.args.get("status", "").strip()
    date_from = request.args.get("date_from", "").strip()
    date_to = request.args.get("date_to", "").strip()

    query = Examinee.query

    if name:
        query = query.filter(Examinee.name.contains(name))
    if department:
        query = query.filter(Examinee.department.contains(department))
    if status in ("pass", "fail"):
        query = query.filter(Examinee.status == status)
    if date_from:
        try:
            dt = datetime.strptime(date_from, "%Y-%m-%d")
            query = query.filter(Examinee.created_at >= dt)
        except ValueError:
            pass
    if date_to:
        try:
            dt = datetime.strptime(date_to, "%Y-%m-%d") + timedelta(days=1)
            query = query.filter(Examinee.created_at < dt)
        except ValueError:
            pass

    query = query.order_by(Examinee.created_at.desc())
    pagination = query.paginate(page=page, per_page=per_page, error_out=False)

    records = []
    for r in pagination.items:
        records.append(
            {
                "id": r.id,
                "name": r.name,
                "department": r.department,
                "score": r.score,
                "total": r.total,
                "percentage": r.percentage,
                "status": r.status,
                "created_at": r.created_at.replace(tzinfo=timezone.utc).isoformat()
                if r.created_at
                else "",
            }
        )

    return jsonify(
        {
            "code": 0,
            "data": {
                "records": records,
                "page": pagination.page,
                "per_page": pagination.per_page,
                "total": pagination.total,
                "pages": pagination.pages,
            },
        }
    )


@app.route("/api/records/<int:record_id>", methods=["GET"])
def get_record_detail(record_id):
    """获取单条考核记录详情"""
    r = db.session.get(Examinee, record_id)
    if not r:
        abort(404)
    return jsonify(
        {
            "code": 0,
            "data": {
                "id": r.id,
                "name": r.name,
                "department": r.department,
                "score": r.score,
                "total": r.total,
                "percentage": r.percentage,
                "status": r.status,
                "answers": json.loads(r.answers_json),
                "created_at": r.created_at.replace(tzinfo=timezone.utc).isoformat()
                if r.created_at
                else "",
            },
        }
    )


# ---------------------------------------------------------------------------
# API - 导出 Excel
# ---------------------------------------------------------------------------

@app.route("/api/records/export", methods=["GET"])
def export_records():
    """导出考核记录为 Excel"""
    name = request.args.get("name", "").strip()
    department = request.args.get("department", "").strip()
    status = request.args.get("status", "").strip()
    date_from = request.args.get("date_from", "").strip()
    date_to = request.args.get("date_to", "").strip()

    query = Examinee.query
    if name:
        query = query.filter(Examinee.name.contains(name))
    if department:
        query = query.filter(Examinee.department.contains(department))
    if status in ("pass", "fail"):
        query = query.filter(Examinee.status == status)
    if date_from:
        try:
            query = query.filter(Examinee.created_at >= datetime.strptime(date_from, "%Y-%m-%d"))
        except ValueError:
            pass
    if date_to:
        try:
            dt = datetime.strptime(date_to, "%Y-%m-%d") + timedelta(days=1)
            query = query.filter(Examinee.created_at < dt)
        except ValueError:
            pass

    query = query.order_by(Examinee.created_at.desc())
    records = query.all()

    wb = Workbook()
    ws = wb.active
    ws.title = "考核记录"

    headers = ["序号", "姓名", "科室", "得分", "总题数", "正确率(%)", "结果", "考试时间"]
    header_font = Font(bold=True)
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for i, r in enumerate(records, 1):
        row = i + 1
        ws.cell(row=row, column=1, value=i)
        ws.cell(row=row, column=2, value=r.name)
        ws.cell(row=row, column=3, value=r.department)
        ws.cell(row=row, column=4, value=r.score)
        ws.cell(row=row, column=5, value=r.total)
        ws.cell(row=row, column=6, value=r.percentage)
        status_label = "合格" if r.status == "pass" else "不合格"
        ws.cell(row=row, column=7, value=status_label)
        ws.cell(
            row=row,
            column=8,
            value=r.created_at.strftime("%Y-%m-%d %H:%M") if r.created_at else "",
        )

    # 调整列宽
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 8
    ws.column_dimensions["E"].width = 8
    ws.column_dimensions["F"].width = 10
    ws.column_dimensions["G"].width = 8
    ws.column_dimensions["H"].width = 18

    export_path = os.path.join(Config.BASE_DIR, "export_temp.xlsx")
    wb.save(export_path)

    return send_file(
        export_path,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f"考核记录_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
    )


# ---------------------------------------------------------------------------
# API - 统计
# ---------------------------------------------------------------------------

@app.route("/api/stats", methods=["GET"])
def get_stats():
    """统计概览"""
    total_examinees = Examinee.query.count()
    if total_examinees == 0:
        return jsonify(
            {
                "code": 0,
                "data": {
                    "total_examinees": 0,
                    "pass_count": 0,
                    "fail_count": 0,
                    "pass_rate": 0,
                    "avg_score": 0,
                },
            }
        )

    pass_count = Examinee.query.filter(Examinee.status == "pass").count()
    fail_count = total_examinees - pass_count
    pass_rate = round(pass_count / total_examinees * 100, 1)

    from sqlalchemy import func

    avg = db.session.query(func.avg(Examinee.percentage)).scalar() or 0

    return jsonify(
        {
            "code": 0,
            "data": {
                "total_examinees": total_examinees,
                "pass_count": pass_count,
                "fail_count": fail_count,
                "pass_rate": pass_rate,
                "avg_score": round(float(avg), 1),
            },
        }
    )


# ---------------------------------------------------------------------------
# API - 角色管理
# ---------------------------------------------------------------------------

@app.route("/api/roles", methods=["GET"])
def list_roles():
    roles = Role.query.order_by(Role.created_at).all()
    items = []
    for r in roles:
        configs = RoleConfig.query.filter_by(role_id=r.id).all()
        items.append(
            {
                "id": r.id,
                "name": r.name,
                "description": r.description,
                "configs": [
                    {
                        "id": c.id,
                        "chapterId": c.chapter_id,
                        "chapterTitle": c.chapter_title,
                        "questionCount": c.question_count,
                    }
                    for c in configs
                ],
            }
        )
    return jsonify({"code": 0, "data": items})


@app.route("/api/roles", methods=["POST"])
def create_role():
    data = request.get_json(force=True)
    name = (data.get("name") or "").strip()
    if not name:
        return jsonify({"code": 1, "msg": "角色名称不能为空"}), 400
    if Role.query.filter_by(name=name).first():
        return jsonify({"code": 1, "msg": "角色已存在"}), 409
    role = Role(name=name, description=data.get("description", ""))
    db.session.add(role)
    db.session.commit()
    return jsonify({"code": 0, "data": {"id": role.id, "name": role.name}})


@app.route("/api/roles/<int:role_id>", methods=["PUT"])
def update_role(role_id):
    role = db.session.get(Role, role_id)
    if not role:
        abort(404)
    data = request.get_json(force=True)
    name = (data.get("name") or "").strip()
    if name:
        role.name = name
    if "description" in data:
        role.description = data.get("description", "")
    db.session.commit()
    return jsonify({"code": 0})


@app.route("/api/roles/<int:role_id>", methods=["DELETE"])
def delete_role(role_id):
    role = db.session.get(Role, role_id)
    if not role:
        abort(404)
    db.session.delete(role)
    db.session.commit()
    return jsonify({"code": 0})


@app.route("/api/roles/<int:role_id>/config", methods=["PUT"])
def save_role_config(role_id):
    role = db.session.get(Role, role_id)
    if not role:
        abort(404)
    data = request.get_json(force=True)
    chapters = data.get("chapters", [])

    # 删除旧配置
    RoleConfig.query.filter_by(role_id=role_id).delete()

    # 写入新配置
    for ch in chapters:
        cfg = RoleConfig(
            role_id=role_id,
            chapter_id=ch["chapterId"],
            chapter_title=ch.get("chapterTitle", ""),
            question_count=ch.get("questionCount", 0),
        )
        db.session.add(cfg)
    db.session.commit()
    return jsonify({"code": 0})


# ---------------------------------------------------------------------------
# API - 考试会话（续答 / 记录）
# ---------------------------------------------------------------------------

@app.route("/api/session/start", methods=["POST"])
def start_session():
    """开始新考试会话（检测到未完成会话时自动续答）"""
    data = request.get_json(force=True)
    name = (data.get("name") or "").strip()
    department = (data.get("department") or "").strip()
    role_id = data.get("roleId")

    if not name or not department:
        return jsonify({"code": 1, "msg": "请填写姓名和科室"}), 400

    # 检测是否存在未完成会话
    existing = ExamSession.query.filter_by(
        name=name, department=department, status="in_progress"
    ).first()
    if existing:
        questions = json.loads(existing.questions_json)
        answers = json.loads(existing.answers_json)
        safe_qs = [
            {
                "id": q["id"],
                "chapterId": q.get("chapterId", 0),
                "chapterTitle": q.get("chapterTitle", ""),
                "type": q["type"],
                "stem": q["stem"],
                "options": q["options"],
                "answer": q.get("answer", []),
            }
            for q in questions
        ]
        return jsonify(
            {
                "code": 0,
                "data": {
                    "sessionId": existing.id,
                    "resume": True,
                    "currentIndex": existing.current_index,
                    "questions": safe_qs,
                    "answers": answers,
                },
            }
        )

    # 获取角色信息
    role = db.session.get(Role, role_id) if role_id else None
    role_name = role.name if role else ""

    # 选取题目
    all_questions = load_questions_from_db_or_file()
    if role:
        configs = RoleConfig.query.filter_by(role_id=role.id).all()
        selected = []
        for cfg in configs:
            chapter_qs = [q for q in all_questions if q.get("chapterId") == cfg.chapter_id]
            if cfg.question_count > 0 and len(chapter_qs) > cfg.question_count:
                chapter_qs = _random.sample(chapter_qs, cfg.question_count)
            selected.extend(chapter_qs)
        if not selected:
            selected = all_questions
    else:
        selected = all_questions

    _random.shuffle(selected)

    session = ExamSession(
        name=name,
        department=department,
        role_id=role_id,
        role_name=role_name,
        questions_json=json.dumps(selected, ensure_ascii=False),
        answers_json=json.dumps(
            [{"selected": [], "revealed": False} for _ in selected], ensure_ascii=False
        ),
        ip_address=request.remote_addr or "",
    )
    db.session.add(session)
    db.session.commit()

    safe_qs = [
        {
            "id": q["id"],
            "chapterId": q.get("chapterId", 0),
            "chapterTitle": q.get("chapterTitle", ""),
            "type": q["type"],
            "stem": q["stem"],
            "options": q["options"],
            "answer": q.get("answer", []),
        }
        for q in selected
    ]
    return jsonify(
        {
            "code": 0,
            "data": {
                "sessionId": session.id,
                "resume": False,
                "currentIndex": 0,
                "questions": safe_qs,
                "answers": json.loads(session.answers_json),
            },
        }
    )


@app.route("/api/session/<int:session_id>", methods=["GET"])
def get_session(session_id):
    session = db.session.get(ExamSession, session_id)
    if not session:
        abort(404)
    questions = json.loads(session.questions_json)
    answers = json.loads(session.answers_json)
    safe_qs = [
        {
            "id": q["id"],
            "chapterId": q.get("chapterId", 0),
            "chapterTitle": q.get("chapterTitle", ""),
            "type": q["type"],
            "stem": q["stem"],
            "options": q["options"],
            "answer": q.get("answer", []),
        }
        for q in questions
    ]
    return jsonify(
        {
            "code": 0,
            "data": {
                "sessionId": session.id,
                "status": session.status,
                "currentIndex": session.current_index,
                "questions": safe_qs,
                "answers": answers,
                "roleName": session.role_name,
            },
        }
    )


@app.route("/api/session/<int:session_id>/answer", methods=["PUT"])
def update_session_answer(session_id):
    """保存单题答案（自动保存）"""
    session = db.session.get(ExamSession, session_id)
    if not session or session.status != "in_progress":
        return jsonify({"code": 1, "msg": "会话不存在或已结束"}), 404

    data = request.get_json(force=True)
    question_index = data.get("questionIndex")
    selected = data.get("selected", [])
    revealed = data.get("revealed", True)
    current_index = data.get("currentIndex")

    answers = json.loads(session.answers_json)
    if question_index is not None and 0 <= question_index < len(answers):
        answers[question_index] = {"selected": selected, "revealed": revealed}
    if current_index is not None:
        session.current_index = current_index

    session.answers_json = json.dumps(answers, ensure_ascii=False)
    db.session.commit()
    return jsonify({"code": 0})


@app.route("/api/session/<int:session_id>/finish", methods=["POST"])
def finish_session(session_id):
    """完成考试并计分"""
    session = db.session.get(ExamSession, session_id)
    if not session or session.status != "in_progress":
        return jsonify({"code": 1, "msg": "会话不存在或已结束"}), 404

    questions = json.loads(session.questions_json)
    answers = json.loads(session.answers_json)
    q_map = {q["id"]: q for q in questions}

    correct_count = 0
    detail = []
    for i, qdata in enumerate(questions):
        q = q_map.get(qdata["id"])
        if not q:
            continue
        ans = answers[i] if i < len(answers) else {"selected": []}
        selected = ans.get("selected", [])
        correct_set = set(q["answer"])
        selected_set = set(selected)
        is_correct = correct_set == selected_set
        if is_correct:
            correct_count += 1
        detail.append(
            {
                "questionId": q["id"],
                "selected": selected,
                "correct": is_correct,
                "correctAnswer": q["answer"],
            }
        )

    total = len(questions)
    percentage = round(correct_count / total * 100, 1) if total > 0 else 0
    status_result = "pass" if percentage >= app.config["PASS_THRESHOLD"] else "fail"

    session.score = correct_count
    session.total = total
    session.percentage = percentage
    session.status_result = status_result
    session.status = "finished"
    session.finished_at = datetime.now(timezone.utc)
    db.session.commit()

    # 同步写入 Examinee 表（管理后台向后兼容）
    examinee = Examinee(
        name=session.name,
        department=session.department,
        score=correct_count,
        total=total,
        percentage=percentage,
        status=status_result,
        answers_json=json.dumps(detail, ensure_ascii=False),
        ip_address=session.ip_address or "",
    )
    db.session.add(examinee)
    db.session.commit()

    return jsonify(
        {
            "code": 0,
            "data": {
                "id": examinee.id,
                "sessionId": session.id,
                "score": correct_count,
                "total": total,
                "percentage": percentage,
                "status": status_result,
            },
        }
    )


@app.route("/api/session/<int:session_id>/result", methods=["GET"])
def get_session_result(session_id):
    """获取已完成考试的结果"""
    session = db.session.get(ExamSession, session_id)
    if not session or session.status != "finished":
        return jsonify({"code": 1, "msg": "考试未完成或不存在"}), 404

    questions = json.loads(session.questions_json)
    answers = json.loads(session.answers_json)

    detail = []
    for i, q in enumerate(questions):
        ans = answers[i] if i < len(answers) else {"selected": []}
        detail.append(
            {
                "questionId": q["id"],
                "stem": q.get("stem", ""),
                "type": q.get("type", ""),
                "options": q.get("options", []),
                "selected": ans.get("selected", []),
                "correctAnswer": q.get("answer", []),
                "correct": set(ans.get("selected", [])) == set(q.get("answer", [])),
            }
        )

    return jsonify(
        {
            "code": 0,
            "data": {
                "sessionId": session.id,
                "name": session.name,
                "department": session.department,
                "roleName": session.role_name,
                "score": session.score,
                "total": session.total,
                "percentage": session.percentage,
                "status": session.status_result,
                "finished_at": session.finished_at.replace(tzinfo=timezone.utc).isoformat()
                if session.finished_at
                else "",
                "details": detail,
            },
        }
    )


# ---------------------------------------------------------------------------
# API - 按章节分组题目
# ---------------------------------------------------------------------------

@app.route("/api/questions/chapters", methods=["GET"])
def get_questions_by_chapter():
    """返回按章节分组的题目（含答案，供题库浏览使用）"""
    questions = load_questions_from_db_or_file()
    chapters = {}
    for q in questions:
        cid = q.get("chapterId", 0)
        if cid not in chapters:
            chapters[cid] = {
                "chapterId": cid,
                "chapterTitle": q.get("chapterTitle", f"第{cid}章"),
                "questions": [],
            }
        chapters[cid]["questions"].append(
            {
                "id": q["id"],
                "type": q["type"],
                "stem": q["stem"],
                "options": q["options"],
                "answer": q.get("answer", []),
            }
        )
    sorted_chapters = sorted(chapters.values(), key=lambda x: x["chapterId"])
    return jsonify({"code": 0, "data": sorted_chapters})


# ---------------------------------------------------------------------------
# API - 个人历史记录
# ---------------------------------------------------------------------------

@app.route("/api/my-records", methods=["GET"])
def get_my_records():
    name = request.args.get("name", "").strip()
    department = request.args.get("department", "").strip()
    if not name or not department:
        return jsonify({"code": 1, "msg": "缺少姓名或科室"}), 400

    sessions = (
        ExamSession.query.filter_by(name=name, department=department, status="finished")
        .order_by(ExamSession.finished_at.desc())
        .all()
    )
    records = []
    for s in sessions:
        records.append(
            {
                "id": s.id,
                "roleName": s.role_name,
                "score": s.score,
                "total": s.total,
                "percentage": s.percentage,
                "status": s.status_result,
                "finished_at": s.finished_at.replace(tzinfo=timezone.utc).isoformat()
                if s.finished_at
                else "",
            }
        )
    return jsonify({"code": 0, "data": records})


# ---------------------------------------------------------------------------
# 初始化数据库
# ---------------------------------------------------------------------------

def init_database():
    """创建表并导入初始题库"""
    with app.app_context():
        db.create_all()
        # 检查是否已有试卷
        if ExamPaper.query.first() is None:
            q_path = os.path.join(Config.BASE_DIR, "static", "js", "questions.json")
            if os.path.exists(q_path):
                with open(q_path, "r", encoding="utf-8") as f:
                    questions = json.load(f)
                paper = ExamPaper(
                    version="1.0",
                    description="初始题库",
                    questions_json=json.dumps(questions, ensure_ascii=False),
                    total_questions=len(questions),
                )
                db.session.add(paper)
                db.session.commit()
                print(f"题库导入完成：共 {len(questions)} 题")
            else:
                print("警告：static/js/questions.json 不存在，题库未导入")


# ---------------------------------------------------------------------------
# 入口
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    init_database()
    app.run(host="0.0.0.0", port=5000, debug=True)
